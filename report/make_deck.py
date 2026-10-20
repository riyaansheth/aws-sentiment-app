"""Build the presentation deck. Run: python3 make_deck.py (any cwd)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image as PI

OUT = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(OUT); SS = os.path.join(ROOT, "screenshots")
INK, MIST, DIM, AMBER, COBALT, SAGE = "0B1F26", "A9C7CC", "5F7F86", "FFC65C", "4D6BFF", "8FB8AE"
DISPLAY, BODY = "Avenir Next Heavy", "Avenir Next"
W, H = Inches(13.333), Inches(7.5)
prs = Presentation(); prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

def rgb(h): return RGBColor.from_string(h)
def bg(slide, color=INK):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); r.fill.solid(); r.fill.fore_color.rgb = rgb(color); r.line.fill.background(); return r
def text(slide, x, y, w, h, runs, size=18, color=MIST, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for item in (runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.line_spacing = spacing
        if isinstance(item, tuple): t, kw = item
        else: t, kw = item, {}
        r = p.add_run(); r.text = t; r.font.name = kw.get("font", font); r.font.size = Pt(kw.get("size", size)); r.font.color.rgb = rgb(kw.get("color", color)); r.font.bold = kw.get("bold", False)
        p.space_after = Pt(kw.get("after", 6))
    return tb
def pic(slide, path, x, y, w=None, h=None):
    iw, ih = PI.open(path).size
    if w and not h: h = int(w * ih / iw)
    if h and not w: w = int(h * iw / ih)
    return slide.shapes.add_picture(path, x, y, w, h)
def full_bleed(slide, path, dim=0):
    p = pic(slide, path, 0, 0, w=W)
    if p.height > H: p.crop_bottom = (p.height - H) / p.height; p.height = H
    elif p.height < H: p.height = H; p.width = int(H * PI.open(path).size[0] / PI.open(path).size[1]); p.left = int((W - p.width) / 2)
    if dim:
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); r.fill.solid(); r.fill.fore_color.rgb = rgb(INK); r.fill.fore_color.brightness = 0
        from pptx.oxml.ns import qn
        sf = r.fill._xPr.find(qn("a:solidFill")); clr = sf[0]; a = clr.makeelement(qn("a:alpha"), {"val": str(int(dim*1000))}); clr.append(a)
        r.line.fill.background()
def headline(slide, word, color=MIST, size=96, y=Inches(2.2), x=Inches(0.8), w=Inches(7)):
    return text(slide, x, y, w, Inches(3), [(word, {"font": DISPLAY, "size": size, "color": color})], spacing=0.9)
def footer(slide, n, color=DIM):
    text(slide, Inches(0.8), H - Inches(0.6), Inches(6), Inches(0.3), [("Mood Orb  |  Cloud Application Development  |  Riyaan Sheth", {"size": 10, "color": color})])
    text(slide, W - Inches(1.6), H - Inches(0.6), Inches(0.8), Inches(0.3), [(str(n), {"size": 10, "color": color})], align=PP_ALIGN.RIGHT)
def bullets(slide, items, x=Inches(7.2), y=Inches(1.6), w=Inches(5.4), size=18, color=MIST, gap=14):
    runs = []
    for it in items:
        if isinstance(it, tuple): runs.append((it[0], {"size": size, "color": it[1] if len(it) > 1 else color, "after": gap, "font": DISPLAY if len(it) > 2 else BODY}))
        else: runs.append((it, {"size": size, "color": color, "after": gap}))
    return text(slide, x, y, w, Inches(5.2), runs)

n = 0
def slide(): 
    global n; n += 1; return prs.slides.add_slide(BLANK), n

# 1 title
s, i = slide(); full_bleed(s, os.path.join(SS, "08a-app-landing.png"), dim=55)
text(s, Inches(0.8), Inches(0.9), Inches(11), Inches(2), [("Mood Orb", {"font": DISPLAY, "size": 110, "color": MIST})], spacing=0.9)
text(s, Inches(0.8), Inches(3.1), Inches(9), Inches(1.4), [("A serverless sentiment analysis web app on AWS", {"size": 26, "color": MIST}), ("Lambda, API Gateway, DynamoDB and S3", {"size": 20, "color": AMBER})])
text(s, Inches(0.8), Inches(5.6), Inches(9), Inches(1.4), [("Riyaan Sheth, App ID 2409365", {"size": 16, "color": MIST}), ("Cloud Application Development (CAP 303), TY B.Tech Semester V", {"size": 14, "color": DIM}), ("Supervisor: Dr. Kunal Meher", {"size": 14, "color": DIM})])

# 2 what it does
s, i = slide(); bg(s); headline(s, "Type\nanything.", size=88)
bullets(s, ["You write a sentence. The app reads its mood and answers Positive, Negative, Neutral or Mixed with a score for each.",
            "A 3D orb on the page changes colour and motion to match the answer.",
            "Every reading is saved, and a History panel brings them back.",
            ("Live at 9b0jyd6730.execute-api.ap-southeast-2.amazonaws.com", AMBER)], y=Inches(2.3)); footer(s, i)

# 3 architecture
s, i = slide(); bg(s, "FFFFFF"); text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(1), [("Four services, one request path", {"font": DISPLAY, "size": 40, "color": INK})])
pic(s, os.path.join(OUT, "block_diagram.png"), Inches(0.8), Inches(1.6), w=Inches(11.7)); footer(s, i, DIM)

# 4 services
s, i = slide(); bg(s); headline(s, "Who does\nwhat", size=80, y=Inches(0.7))
items = [("API Gateway", "The public HTTPS door. Three routes, all pointing at one function."),
         ("Lambda", "The only compute. Serves the page, scores the text, writes the results."),
         ("DynamoDB", "One item per reading. History reads from here."),
         ("S3", "One JSON object per reading. A plain file archive of everything scored.")]
y = Inches(0.9)
for name, desc in items:
    text(s, Inches(6.6), y, Inches(6), Inches(0.6), [(name, {"font": DISPLAY, "size": 30, "color": AMBER})]); text(s, Inches(6.6), y + Inches(0.6), Inches(6), Inches(0.9), [(desc, {"size": 17})]); y += Inches(1.55)
text(s, Inches(0.8), Inches(5.4), Inches(5.5), Inches(1), [("Plus IAM for the execution role and CloudWatch for logs. Region: Sydney. Cost: nothing, inside the Free Plan.", {"size": 14, "color": DIM})]); footer(s, i)

# 5 request flow (a real sequence, so numbering is earned)
s, i = slide(); bg(s); headline(s, "One click,\nfour hops", size=80, y=Inches(0.7))
steps = [("1", "Browser sends POST /analyze with the text."), ("2", "API Gateway proxies the request into Lambda."), ("3", "Lambda scores the text, writes S3 and DynamoDB."), ("4", "JSON comes back and the orb changes mood.")]
y = Inches(1.0)
for num, t in steps:
    text(s, Inches(6.6), y, Inches(0.8), Inches(0.9), [(num, {"font": DISPLAY, "size": 44, "color": COBALT})]); text(s, Inches(7.5), y + Inches(0.15), Inches(5.2), Inches(0.9), [(t, {"size": 19})]); y += Inches(1.25)
text(s, Inches(6.6), Inches(6.0), Inches(6), Inches(0.6), [("GET / returns the page itself. GET /history returns the last 20 items.", {"size": 14, "color": DIM})]); footer(s, i)

# 6 scoring
s, i = slide(); bg(s); headline(s, "How it\nscores", size=80, y=Inches(0.7))
bullets(s, ["Two word lists: positive words and negative words.",
            "A word directly after not, never, don't or can't flips polarity. \"not good\" counts as negative.",
            "Label comes from the counts: more positive wins, more negative wins, no hits is Neutral, a close split is Mixed.",
            "Scores are the share of each count, so the app can draw bars.",
            ("\"I love this, it is great\"  ->  Positive 100", AMBER), ("\"not good at all\"  ->  Negative 100", COBALT)], y=Inches(1.0), size=17, gap=12); footer(s, i)

# 7 why not comprehend
s, i = slide(); bg(s, COBALT); headline(s, "Blocked.", color="FFFFFF", size=120, y=Inches(1.4))
bullets(s, [("The plan was Amazon Comprehend.", "FFFFFF"),
            ("The account is on the AWS Free Plan. Its organisation policy explicitly denies comprehend:DetectSentiment, and locks the account to one region.", "FFFFFF"),
            ("First real call: AccessDeniedException, with the policy ARN in the message.", "FFFFFF"),
            ("So scoring moved inside Lambda. The call site is one line, so Comprehend drops back in on a paid account.", "FFFFFF")], y=Inches(1.5), size=18); footer(s, i, "DCE2FF")

# 8 deployment shots
s, i = slide(); bg(s); text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8), [("Built in the console, step by step", {"font": DISPLAY, "size": 36})])
grid = ["01-s3-bucket-created.png", "02-lambda-function-overview.png", "04-api-gateway-routes.png", "05-dynamodb-table.png"]
caps = ["S3 bucket", "Lambda function with API trigger", "API Gateway routes", "DynamoDB table"]
for k, (fn, c) in enumerate(zip(grid, caps)):
    x = Inches(0.8) + (k % 2) * Inches(6.1); y = Inches(1.65) + (k // 2) * Inches(2.95)
    text(s, x, y - Inches(0.32), Inches(5.6), Inches(0.3), [(c, {"size": 12, "color": DIM})]); pic(s, os.path.join(SS, fn), x, y, w=Inches(4.9))

# 9 stored results
s, i = slide(); bg(s); text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8), [("Every reading lands in two places", {"font": DISPLAY, "size": 36})])
pic(s, os.path.join(SS, "06-s3-results-objects.png"), Inches(0.8), Inches(1.5), w=Inches(5.7)); text(s, Inches(0.8), Inches(4.85), Inches(5.7), Inches(0.3), [("S3: results/<uuid>.json", {"size": 12, "color": DIM})])
pic(s, os.path.join(SS, "07-dynamodb-items.png"), Inches(6.8), Inches(1.5), w=Inches(5.7)); text(s, Inches(6.8), Inches(4.85), Inches(5.7), Inches(0.3), [("DynamoDB: SentimentResults items", {"size": 12, "color": DIM})])
footer(s, i)

# 10 demo result
s, i = slide(); full_bleed(s, os.path.join(SS, "08-app-live-result.png")); footer(s, i, MIST)
# 11 demo history
s, i = slide(); full_bleed(s, os.path.join(SS, "09-app-history.png")); footer(s, i, MIST)

# 12 limits
s, i = slide(); bg(s); headline(s, "Where it\nbreaks", size=80, y=Inches(0.7), w=Inches(6))
bullets(s, [("\"What a wonderful way to spend my birthday: the hotel lost our booking.\"", AMBER),
            "Scored Positive. A word list cannot read sarcasm, distant negation, or words it has never seen.",
            "Fix: swap the scorer for Amazon Comprehend on a standard account. One line.",
            "Next: Cognito sign in for per user history, CloudFront with a custom domain, a DynamoDB index on time so History is a query, not a scan."], y=Inches(1.0), size=17, gap=14); footer(s, i)

# 13 close
s, i = slide(); bg(s); headline(s, "Thanks.", color=AMBER, size=120, y=Inches(1.6))
text(s, Inches(0.8), Inches(4.3), Inches(11), Inches(2), [("Live: https://9b0jyd6730.execute-api.ap-southeast-2.amazonaws.com/", {"size": 18}), ("Code: https://github.com/riyaansheth/aws-sentiment-app", {"size": 18}), ("Riyaan Sheth, 2409365", {"size": 16, "color": DIM})])

out = os.path.join(OUT, "Mood_Orb_Presentation.pptx"); prs.save(out); print(out, n, "slides")
