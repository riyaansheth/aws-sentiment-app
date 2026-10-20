"""Build the mini-project report as DOCX and PDF. Run: python3 make_report.py (any cwd)."""
import os
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, Preformatted, Table, TableStyle, KeepTogether

OUT = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(OUT)
SS = os.path.join(ROOT, "screenshots")

TITLE = "Sentiment Analysis Web App using AWS Lambda, API Gateway, DynamoDB and S3"
STUDENTS = [("Riyaan Sheth", "2409365")]
SUPERVISOR = "Dr. Kunal Meher"
LIVE = "https://9b0jyd6730.execute-api.ap-southeast-2.amazonaws.com/"
REPO = "https://github.com/riyaansheth/aws-sentiment-app"

INTRO = [
"This project is a serverless web application that reads a piece of text and reports its sentiment as Positive, Negative, Neutral or Mixed, along with a score for each class. Every reading is stored so that it can be listed again later. The application is built entirely on Amazon Web Services and uses four managed services: Amazon API Gateway, AWS Lambda, Amazon DynamoDB and Amazon S3. No servers are provisioned or managed; each component scales on demand and costs nothing at the project's usage level under the AWS Free Plan.",
"The user opens a single URL served by API Gateway. A landing page shows a three dimensional orb rendered in the browser with WebGL. Clicking the orb moves it aside and reveals the analyzer. The user types text and submits it. The browser sends the text to the API, the Lambda function scores it and writes the result to both S3 and DynamoDB, and the response is shown next to the orb, which changes colour and motion to match the mood. A History panel calls a third route that reads the most recent results back from DynamoDB.",
"The objectives of the project were to design a complete request path through several AWS services, to keep the whole system inside the free tier, to expose it on a public HTTPS endpoint without any manual server setup, and to document the deployment with screenshots of each console step.",
"Figure 1 shows the block diagram. API Gateway is the only public entry point. Lambda is the only compute. S3 holds one JSON object per reading and DynamoDB holds one item per reading, which gives both a file archive and a queryable table from the same write.",
]
DEV_TECH = [
("Python 3.14", "Language for the Lambda function. The function has no third party dependencies beyond boto3, which the Lambda runtime provides."),
("boto3", "The AWS SDK for Python, used to put objects into S3 and to put and scan items in DynamoDB."),
("HTML, CSS and JavaScript", "The single page frontend. It is a static file packaged inside the Lambda deployment and returned on GET /."),
("Three.js r128", "WebGL library used to render the orb. A custom vertex shader displaces an icosphere with layered noise and a fragment shader lights and colours it from uniforms that the page updates after each reading."),
("Syne typeface", "Loaded from Google Fonts for the page typography."),
("Git and GitHub", "Version control. The full source is at " + REPO + "."),
("Playwright", "Used during development to drive the AWS console in a browser and capture the deployment screenshots."),
]
DEPLOY_TECH = [
("Amazon API Gateway (HTTP API)", "Provides the public HTTPS endpoint. The API sentiment-api has three routes, GET /, POST /analyze and GET /history, all integrated with the Lambda function through a proxy integration. The $default stage auto-deploys on every change."),
("AWS Lambda", "Runs the function sentimentFn on Python 3.14 with 128 MB memory and a 15 second timeout. The function serves the frontend, scores the text, stores results and reads history. Environment variables BUCKET and TABLE name the storage targets."),
("Amazon DynamoDB", "Table SentimentResults with partition key id (String) in on-demand capacity mode. Each reading is one item holding the text, label, scores and timestamp."),
("Amazon S3", "Bucket sentiment-app-riyaan with public access blocked. Each reading is written as results/<uuid>.json."),
("AWS IAM", "The function executes under the role PowerUserRole, which grants the S3 and DynamoDB permissions it needs."),
("Amazon CloudWatch Logs", "Receives the function's log output automatically for debugging."),
("Region", "All resources are in Asia Pacific (Sydney), ap-southeast-2."),
]
SCORER_NOTE = "The original design used Amazon Comprehend for sentiment detection. The AWS account used for this project is on the AWS Free Plan, which places the account inside an organisation whose service control policy explicitly denies the comprehend:DetectSentiment action. The first test call returned an AccessDeniedException naming that policy. The scoring was therefore implemented inside the Lambda function as a lexicon based classifier: it counts positive and negative words from two word lists, flips a word's polarity when it directly follows a negation such as not or never, and derives the label and class scores from the two counts. The call site is a single line, so Amazon Comprehend can be substituted on a paid account without any other change."
SHOTS = [
("01-s3-bucket-created.png", "Amazon S3 bucket sentiment-app-riyaan created in ap-southeast-2 with public access blocked."),
("02-lambda-function-overview.png", "AWS Lambda function sentimentFn with the API Gateway trigger attached."),
("02b-lambda-code-editor.png", "The function source in the Lambda code editor after uploading the deployment package."),
("03-iam-execution-role.png", "IAM role PowerUserRole used as the function's execution role."),
("04-api-gateway-routes.png", "API Gateway HTTP API sentiment-api with routes GET /, POST /analyze and GET /history, each integrated with the Lambda function."),
("05-dynamodb-table.png", "DynamoDB table SentimentResults with partition key id."),
("06-s3-results-objects.png", "Result objects written by the function under the results/ prefix in S3."),
("07-dynamodb-items.png", "The same results stored as items in DynamoDB."),
("08a-app-landing.png", "Landing page served from the API Gateway URL. Clicking the orb opens the analyzer."),
("08-app-live-result.png", "Analyzer showing a Positive reading with class scores. The orb has changed colour to match."),
("09-app-history.png", "History panel listing recent readings fetched from DynamoDB through GET /history."),
]
CONCLUSION = [
"The project delivers a working, publicly reachable sentiment analysis application built from four AWS services with no servers to manage. API Gateway, Lambda, DynamoDB and S3 each perform one clear role, and the whole request path from browser to storage and back was deployed and verified through the AWS console, with every step captured in the screenshots above.",
"Working inside the AWS Free Plan shaped the design. The region was restricted to Sydney and Amazon Comprehend was blocked by policy, so the sentiment logic moved into the Lambda function as a lexicon classifier. This kept the project free of cost but it also set the main limitation: a word list cannot read sarcasm, distant negation or words outside its vocabulary, so a sentence such as 'what a wonderful way to lose our booking' is scored as Positive. Replacing the scorer with Amazon Comprehend on a standard account is a one line change and would be the first improvement.",
"Other extensions that fit the same architecture are per user history with Amazon Cognito sign in, a CloudFront distribution in front of the API for caching and a custom domain, and a DynamoDB global secondary index on the timestamp so that history can be queried in order instead of scanned.",
"The main learning outcomes were how an HTTP API proxies requests into Lambda, how a single function can serve both a frontend and a JSON API, how IAM roles grant service to service permissions, and how account level policies can constrain which services are available and must be checked early in a design.",
]

def read(p): return open(os.path.join(ROOT, p)).read()
LAMBDA_SRC = read("lambda_function.py")
html = read("index.html")
def section(src, start, end): i = src.index(start); j = src.index(end, i); return src[i:j].rstrip()
HTML_EXCERPTS = [
("Shader that shapes and lights the orb (index.html)", section(html, "const vert = `", "const mat = new THREE.ShaderMaterial")),
("Sending text to the API and showing the result (index.html)", section(html, "function showResult(j){", "t.addEventListener('input'")),
("Landing page transition (index.html)", section(html, "const enter = document.getElementById('enter');", "// ---------- UI ----------")),
]

def build_docx():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    for s in doc.sections: s.top_margin = s.bottom_margin = Cm(2.2); s.left_margin = s.right_margin = Cm(2.5)
    def center(text, size=12, bold=False, space=6):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); p.paragraph_format.space_after = Pt(space); return p
    def h(text, lvl=1):
        p = doc.add_heading(text, level=lvl)
        for r in p.runs: r.font.name = "Times New Roman"
        return p
    def para(text):
        p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(8); return p
    def code(text):
        for line in text.split("\n"):
            p = doc.add_paragraph(); r = p.add_run(line); r.font.name = "Courier New"; r.font.size = Pt(8.5)
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    def fig(path, caption, width=6.2):
        doc.add_picture(path, width=Inches(width)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        center(caption, size=10, space=12)
    def names_table():
        t = doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for n, i in STUDENTS:
            c = t.add_row().cells; c[0].text = n; c[1].text = i
            for cell in c:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs: r.bold = True
    logo = os.path.join(OUT, "logo-000.png")
    doc.add_picture(logo, width=Inches(4.2)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    center("uGDX School of Technology", 18, True, 18)
    center("Mini Project Report on", 12, False, 12)
    center("“" + TITLE + "”", 14, True, 14)
    center("Submitted in partial fulfillment of the requirements", 12); center("for the course", 12, space=14)
    center("Course: Cloud Application Development (CAP 303)", 12, True); center("Class: TY B. Tech (Semester V)", 12, True, 14)
    center("By", 12, space=8); names_table(); doc.add_paragraph()
    center("Supervisor:", 12); center(SUPERVISOR, 12); center("(Assistant Professor)", 12, space=24)
    center("AY: 2026-2027", 12); doc.add_page_break()
    doc.add_picture(logo, width=Inches(4.2)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    center("uGDX School of Technology", 18, True, 18); center("CERTIFICATE", 13, True, 14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("This is to certify that the project entitled “"); p.add_run(TITLE).bold = True; p.add_run("” is a bonafide work of")
    names_table(); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("submitted in partial fulfilment of the requirement for the course “"); p.add_run("Cloud Application Development").bold = True
    p.add_run("” of Third Year of Bachelor degree in Technology (CS, AI & ML) during academic year 2026-27.")
    for _ in range(6): doc.add_paragraph()
    t = doc.add_table(rows=2, cols=2); t.cell(0,0).text = "Supervisor/Guide"; t.cell(0,1).text = "External Examiner"; t.cell(1,0).text = "(" + SUPERVISOR + ")"
    t.cell(0,1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_page_break()
    center("Table of Contents", 16, True, 18)
    for line in ["1. Introduction", "2. Technologies used", "      a. Development Technology used", "      b. Deployment Technology used", "3. Code", "4. Screenshots", "5. Conclusion"]:
        p = doc.add_paragraph(); p.add_run(line).bold = True; p.paragraph_format.space_after = Pt(10)
    doc.add_page_break()
    h("1. Introduction")
    for t_ in INTRO: para(t_)
    fig(os.path.join(OUT, "block_diagram.png"), "Figure 1. Block diagram of the application.", 6.4)
    para("Live application: " + LIVE); para("Source code: " + REPO)
    doc.add_page_break()
    h("2. Technologies used"); h("a. Development Technology used", 2)
    for n, d in DEV_TECH:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(n + ". ").bold = True; p.add_run(d)
    h("b. Deployment Technology used", 2)
    for n, d in DEPLOY_TECH:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(n + ". ").bold = True; p.add_run(d)
    h("Note on the sentiment scorer", 2); para(SCORER_NOTE)
    doc.add_page_break()
    h("3. Code"); para("The backend is a single file, lambda_function.py, reproduced in full below. The frontend, index.html, is about 300 lines; the three parts that matter most follow the backend listing. The complete source is in the GitHub repository.")
    h("lambda_function.py", 2); code(LAMBDA_SRC)
    for title, src in HTML_EXCERPTS: h(title, 2); code(src)
    doc.add_page_break()
    h("4. Screenshots")
    for i, (fn, cap) in enumerate(SHOTS, 1): fig(os.path.join(SS, fn), f"Figure {i+1}. {cap}", 6.4)
    doc.add_page_break()
    h("5. Conclusion")
    for t_ in CONCLUSION: para(t_)
    out = os.path.join(OUT, "Mini_Project_Report.docx"); doc.save(out); return out

def build_pdf():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman", fontSize=12, leading=16, alignment=TA_JUSTIFY, spaceAfter=8)
    cen = ParagraphStyle("cen", parent=body, alignment=TA_CENTER, spaceAfter=6)
    cenb = ParagraphStyle("cenb", parent=cen, fontName="Times-Bold")
    h1 = ParagraphStyle("h1", parent=body, fontName="Times-Bold", fontSize=16, leading=20, spaceBefore=6, spaceAfter=10, alignment=0)
    h2 = ParagraphStyle("h2", parent=body, fontName="Times-Bold", fontSize=13, leading=17, spaceBefore=8, spaceAfter=6, alignment=0)
    cap = ParagraphStyle("cap", parent=cen, fontSize=10, leading=13, spaceAfter=14)
    codes = ParagraphStyle("code", fontName="Courier", fontSize=7.4, leading=9)
    bullet = ParagraphStyle("bul", parent=body, leftIndent=14, bulletIndent=2)
    W, H = A4; maxw = W - 50*mm
    from PIL import Image as PI
    def img(path, w=maxw):
        iw, ih = PI.open(path).size; r = RLImage(path, width=w, height=w*ih/iw); r.hAlign = "CENTER"; return r
    def esc(s): return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    def names():
        t = Table([[n, i] for n, i in STUDENTS], colWidths=[70*mm, 40*mm]); t.setStyle(TableStyle([("FONT", (0,0), (-1,-1), "Times-Bold", 12), ("ALIGN", (0,0), (-1,-1), "CENTER")])); return t
    logo = os.path.join(OUT, "logo-000.png")
    S = [Spacer(1, 10*mm), img(logo, 105*mm), Spacer(1, 6*mm), Paragraph("uGDX School of Technology", ParagraphStyle("t", parent=cenb, fontSize=18, leading=22, spaceAfter=18)),
         Paragraph("Mini Project Report on", cen), Spacer(1, 4*mm), Paragraph("“" + esc(TITLE) + "”", ParagraphStyle("tt", parent=cenb, fontSize=14, leading=18, spaceAfter=14)),
         Paragraph("Submitted in partial fulfillment of the requirements", cen), Paragraph("for the course", cen), Spacer(1, 6*mm),
         Paragraph("Course: Cloud Application Development (CAP 303)", cenb), Paragraph("Class: TY B. Tech (Semester V)", cenb), Spacer(1, 6*mm),
         Paragraph("By", cen), Spacer(1, 3*mm), names(), Spacer(1, 8*mm), Paragraph("Supervisor:", cen), Paragraph(SUPERVISOR, cen), Paragraph("(Assistant Professor)", cen),
         Spacer(1, 12*mm), Paragraph("AY: 2026-2027", cen), PageBreak()]
    sig = Table([["Supervisor/Guide", "External Examiner"], ["(" + SUPERVISOR + ")", ""]], colWidths=[maxw/2, maxw/2], style=TableStyle([("FONT", (0,0), (-1,-1), "Times-Roman", 12), ("ALIGN", (1,0), (1,-1), "RIGHT")]))
    S += [Spacer(1, 10*mm), img(logo, 105*mm), Spacer(1, 6*mm), Paragraph("uGDX School of Technology", ParagraphStyle("t2", parent=cenb, fontSize=18, leading=22, spaceAfter=18)),
          Paragraph("CERTIFICATE", ParagraphStyle("c", parent=cenb, fontSize=13, spaceAfter=14)),
          Paragraph("This is to certify that the project entitled “<b>" + esc(TITLE) + "</b>” is a bonafide work of", body), Spacer(1, 3*mm), names(), Spacer(1, 3*mm),
          Paragraph("submitted in partial fulfilment of the requirement for the course “<b>Cloud Application Development</b>” of Third Year of Bachelor degree in Technology (CS, AI &amp; ML) during academic year 2026-27.", body),
          Spacer(1, 50*mm), sig, PageBreak()]
    S += [Spacer(1, 15*mm), Paragraph("Table of Contents", ParagraphStyle("toc", parent=cenb, fontSize=16, spaceAfter=18))]
    for line in ["1. Introduction", "2. Technologies used", "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;a. Development Technology used", "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;b. Deployment Technology used", "3. Code", "4. Screenshots", "5. Conclusion"]:
        S.append(Paragraph("<b>" + line + "</b>", ParagraphStyle("tl", parent=body, leftIndent=40*mm, spaceAfter=10, alignment=0)))
    S.append(PageBreak())
    S.append(Paragraph("1. Introduction", h1))
    for t_ in INTRO: S.append(Paragraph(esc(t_), body))
    S += [img(os.path.join(OUT, "block_diagram.png")), Paragraph("Figure 1. Block diagram of the application.", cap), Paragraph("Live application: " + LIVE, body), Paragraph("Source code: " + REPO, body), PageBreak()]
    S += [Paragraph("2. Technologies used", h1), Paragraph("a. Development Technology used", h2)]
    for n, d in DEV_TECH: S.append(Paragraph("<b>" + esc(n) + ".</b> " + esc(d), bullet, bulletText="•"))
    S.append(Paragraph("b. Deployment Technology used", h2))
    for n, d in DEPLOY_TECH: S.append(Paragraph("<b>" + esc(n) + ".</b> " + esc(d), bullet, bulletText="•"))
    S += [Paragraph("Note on the sentiment scorer", h2), Paragraph(esc(SCORER_NOTE), body), PageBreak()]
    S += [Paragraph("3. Code", h1), Paragraph("The backend is a single file, lambda_function.py, reproduced in full below. The frontend, index.html, is about 300 lines; the three parts that matter most follow the backend listing. The complete source is in the GitHub repository.", body),
          Paragraph("lambda_function.py", h2), Preformatted(LAMBDA_SRC, codes, maxLineLength=96)]
    for title, src in HTML_EXCERPTS: S += [Paragraph(esc(title), h2), Preformatted(src, codes, maxLineLength=96)]
    S.append(PageBreak())
    S.append(Paragraph("4. Screenshots", h1))
    for i, (fn, c) in enumerate(SHOTS, 1): S.append(KeepTogether([img(os.path.join(SS, fn)), Paragraph(f"Figure {i+1}. {esc(c)}", cap)]))
    S.append(PageBreak())
    S.append(Paragraph("5. Conclusion", h1))
    for t_ in CONCLUSION: S.append(Paragraph(esc(t_), body))
    out = os.path.join(OUT, "Mini_Project_Report.pdf")
    def footer(c, d):
        if d.page > 3: c.setFont("Times-Roman", 10); c.drawCentredString(W/2, 12*mm, str(d.page))
    SimpleDocTemplate(out, pagesize=A4, leftMargin=25*mm, rightMargin=25*mm, topMargin=22*mm, bottomMargin=22*mm, title=TITLE, author=STUDENTS[0][0]).build(S, onFirstPage=footer, onLaterPages=footer)
    return out

if __name__ == "__main__":
    for s in INTRO + CONCLUSION + [SCORER_NOTE] + [d for _, d in DEV_TECH + DEPLOY_TECH] + [c for _, c in SHOTS]:
        assert "—" not in s and "–" not in s, "dash found"
    print(build_docx()); print(build_pdf())
